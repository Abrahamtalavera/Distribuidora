"""
Cambio #21: consulta del equipo de ventas.

Dos consultas, solo de lectura:
  - Buscar una factura (por número o por cliente) y ver todo lo que le pasó:
    recorrido, entrega, cobro, productos devueltos y lo que recibió bodega.
  - Listar facturas por ruta, rango de fechas, estado y vendedor, con un
    resumen arriba y descarga en Excel.

Cualquier vendedor puede consultar cualquier factura. Igual que el resto de
módulos, la lógica vive aquí y no en las vistas para poder probarla.
"""

from collections import OrderedDict
from datetime import timedelta
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from django.db.models import Prefetch, Q
from django.utils import timezone

from core.cierre import calcular_resultado
from core.models import EntregaDetalle, Factura

CERO = Decimal("0")
CENTAVO = Decimal("0.01")
PERMISO_VENTAS = "core.consultar_ventas"
MAX_RESULTADOS_BUSQUEDA = 50
MAX_FILAS_LISTADO = 2000
MAX_DIAS_LISTADO = 366
DIAS_POR_DEFECTO = 7

# Estado que ve el vendedor: (nombre, clase de color). Las clases son las
# mismas de las demás pantallas (badge-completa, badge-parcial, ...).
ESTADOS = OrderedDict(
    [
        ("POR_ENTREGAR", ("Por entregar", "pendiente")),
        ("EN_RUTA", ("En ruta", "ruta")),
        ("SIN_VISITAR", ("Sin visitar", "pendiente")),
        ("PARCIAL", ("Entrega parcial", "parcial")),
        ("ENTREGADA", ("Entregada", "completa")),
        ("CON_DEVOLUCION", ("Con devolución", "parcial")),
        ("NO_ENTREGADA", ("No entregada", "noentregada")),
        ("ANULADA", ("Anulada", "pendiente")),
    ]
)

# Opciones del filtro "Estado de entrega" del listado.
FILTROS_ESTADO = OrderedDict(
    [
        ("", ("Todas", None)),
        ("PENDIENTES", ("Pendientes de entregar", {"POR_ENTREGAR", "EN_RUTA", "SIN_VISITAR", "PARCIAL"})),
        ("EN_RUTA", ("En ruta", {"EN_RUTA"})),
        ("ENTREGADA", ("Entregadas", {"ENTREGADA"})),
        ("CON_DEVOLUCION", ("Con devolución", {"CON_DEVOLUCION"})),
        ("NO_ENTREGADA", ("No entregadas", {"NO_ENTREGADA"})),
        ("ANULADA", ("Anuladas", {"ANULADA"})),
    ]
)


def puede_consultar(usuario):
    return bool(
        usuario
        and usuario.is_authenticated
        and usuario.is_active
        and usuario.has_perm(PERMISO_VENTAS)
    )


def _r(valor):
    return (valor or CERO).quantize(CENTAVO, rounding=ROUND_HALF_UP)


def _base():
    return Factura.objects.select_related(
        "cliente", "vendedor", "modalidad_pago", "carga__ruta", "carga__conductor",
        "carga__vehiculo", "cartera",
    ).prefetch_related("lineas")


def anotar(factura):
    """
    Le pone a la factura lo que muestran las pantallas de ventas:
    .estado (código), .estado_nombre, .estado_css, .valor_entregado,
    .valor_devuelto y .saldo. Necesita las líneas (prefetch) y la carga.
    """
    lineas = list(factura.lineas.all())
    resultado, valor_ent, valor_dev = calcular_resultado(lineas)
    carga = factura.carga
    if factura.estado_factura == "ANULADA":
        estado = "ANULADA"
    elif resultado == "SIN_VISITAR":
        if carga is None or carga.estado == "PLANEADA":
            estado = "POR_ENTREGAR"
        elif carga.estado == "EN_RUTA":
            estado = "EN_RUTA"
        else:
            estado = "SIN_VISITAR"
    elif resultado == "COMPLETA":
        estado = "ENTREGADA"
    else:
        estado = resultado  # PARCIAL, CON_DEVOLUCION, NO_ENTREGADA

    factura.estado = estado
    factura.estado_nombre, factura.estado_css = ESTADOS[estado]
    factura.valor_entregado = valor_ent
    factura.valor_devuelto = valor_dev

    # Saldo: el de la cartera si ya existe (se recalcula con cada visita);
    # si no, nadie ha pagado nada: debe el total menos lo devuelto.
    cartera = getattr(factura, "cartera", None) if _tiene_cartera(factura) else None
    if estado == "ANULADA":
        factura.saldo = None
    elif cartera is not None:
        factura.saldo = _r(cartera.saldo_pendiente)
    else:
        factura.saldo = _r(max(CERO, factura.total - valor_dev))
    return factura


def _tiene_cartera(factura):
    try:
        factura.cartera
    except Factura.cartera.RelatedObjectDoesNotExist:
        return False
    return True


# ---------------------------------------------------------------------------
# Buscar una factura
# ---------------------------------------------------------------------------


def buscar_facturas(texto):
    """
    Devuelve (exacta, resultados): `exacta` es la factura cuyo número
    coincide exactamente (o None); `resultados`, las que coinciden en número
    o cliente (nombre o código), las más recientes primero.
    """
    texto = (texto or "").strip()
    if not texto:
        return None, []
    exacta = _base().filter(numero_factura__iexact=texto).first()
    if exacta:
        return anotar(exacta), [exacta]
    resultados = list(
        _base()
        .filter(
            Q(numero_factura__icontains=texto)
            | Q(cliente__nombre__icontains=texto)
            | Q(cliente__codigo__iexact=texto)
        )
        .order_by("-fecha_emision", "-numero_factura")[:MAX_RESULTADOS_BUSQUEDA]
    )
    for f in resultados:
        anotar(f)
    return None, resultados


def obtener_factura(factura_id):
    factura = _base().filter(pk=factura_id).first()
    return anotar(factura) if factura else None


def _fecha_texto(cuando):
    if not cuando:
        return ""
    if hasattr(cuando, "hour"):
        return timezone.localtime(cuando).strftime("%d/%m %H:%M")
    return cuando.strftime("%d/%m/%Y")


def _paso(titulo, cuando=None, detalle="", hecho=None, alerta=False):
    return {
        "titulo": titulo,
        "cuando": _fecha_texto(cuando),
        "detalle": detalle,
        "hecho": bool(cuando) if hecho is None else hecho,
        "alerta": alerta,
    }


def detalle_factura(factura):
    """Todo lo que muestra la pantalla de detalle de una factura."""
    entregas = list(
        factura.entregas.filter(estado="CONFIRMADA")
        .select_related(
            "carga__ruta", "carga__vehiculo", "carga__conductor", "repartidor",
            "motivo_devolucion_general__area_responsable",
        )
        .prefetch_related(
            Prefetch(
                "lineas",
                queryset=EntregaDetalle.objects.select_related(
                    "motivo_devolucion__area_responsable", "factura_detalle"
                ),
            )
        )
        .order_by("fecha_entrega", "id")
    )
    # Cambio #22: llegada al cliente e indicador de GPS de cada visita.
    from core.models import ParadaCliente
    from core.tiempos import indicador_gps

    llegadas = {
        p.carga_id: p.llegada_en
        for p in ParadaCliente.objects.filter(cliente=factura.cliente, carga_id__in={e.carga_id for e in entregas})
    }
    for e in entregas:
        e.llegada = llegadas.get(e.carga_id)
        e.gps = indicador_gps(e, factura.cliente) if e.gps_estado else None
    ultima = entregas[-1] if entregas else None
    carga = factura.carga or (ultima.carga if ultima else None)

    # --- Recorrido ---------------------------------------------------------
    pasos = [_paso("Emitida", factura.fecha_emision)]
    if carga:
        pasos.append(_paso("En carga", carga.fecha_planeada, carga.codigo_reporte))
    else:
        pasos.append(_paso("En carga", detalle="Sin carga asignada"))
    pasos.append(_paso("Cargada al camión", factura.cargada_en, factura.cargada_por))
    pasos.append(_paso("Salió a ruta", carga.salio_a_ruta_en if carga else None))
    if ultima:
        detalle = factura.estado_nombre.lower() if factura.estado in ("CON_DEVOLUCION", "NO_ENTREGADA", "PARCIAL") else ""
        if len(entregas) > 1:
            detalle = f"{len(entregas)} visitas" + (f" · {detalle}" if detalle else "")
        pasos.append(_paso("Visitada", ultima.fecha_entrega, detalle,
                           alerta=factura.estado in ("CON_DEVOLUCION", "NO_ENTREGADA", "PARCIAL")))
    else:
        pasos.append(_paso("Visitada"))
    pasos.append(_paso("Carga cerrada", carga.cerrada_en if carga else None))
    # Un paso sin fecha después de uno con fecha (por datos viejos) se
    # muestra como hecho si alguno posterior ya ocurrió.
    for i, p in enumerate(pasos):
        if not p["hecho"] and any(q["hecho"] for q in pasos[i + 1:]) and p["titulo"] != "En carga":
            p["hecho"] = True

    # --- Productos -----------------------------------------------------------
    por_linea = {}
    for e in entregas:
        for d in e.lineas.all():
            info = por_linea.setdefault(d.factura_detalle_id, {"motivos": [], "inventario": CERO, "merma": CERO, "por_recibir": False})
            if d.cantidad_devuelta > 0:
                motivo = d.motivo_devolucion or e.motivo_devolucion_general
                if motivo:
                    texto = motivo.nombre
                    if motivo.area_responsable:
                        texto += f" ({motivo.area_responsable.nombre})"
                    if texto not in info["motivos"]:
                        info["motivos"].append(texto)
                if d.recepcion_registrada:
                    info["inventario"] += d.recibido_inventario or CERO
                    info["merma"] += d.recibido_merma or CERO
                else:
                    info["por_recibir"] = True
    productos = []
    for linea in factura.lineas.select_related("producto").order_by("id"):
        info = por_linea.get(linea.id, {"motivos": [], "inventario": CERO, "merma": CERO, "por_recibir": False})
        if linea.cantidad_devuelta > 0:
            partes = []
            if info["inventario"]:
                partes.append(f"{_cant(info['inventario'])} a inventario")
            if info["merma"]:
                partes.append(f"{_cant(info['merma'])} merma")
            if info["por_recibir"]:
                partes.append("pendiente de recibir")
            bodega = " · ".join(partes) or "—"
        else:
            bodega = "—"
        productos.append(
            {
                "codigo": linea.producto.codigo,
                "nombre": linea.producto.nombre,
                "unidad": linea.producto.unidad_medida,
                "facturado": linea.cantidad_facturada,
                "entregado": linea.cantidad_entregada,
                "devuelto": linea.cantidad_devuelta,
                "pendiente": linea.pendiente_entrega,
                "motivo": "; ".join(info["motivos"]) or ("—" if not linea.cantidad_devuelta else "Sin motivo"),
                "bodega": bodega,
                "por_recibir": info["por_recibir"],
            }
        )

    # --- Cobro ---------------------------------------------------------------
    pagos = []
    if _tiene_cartera(factura):
        pagos = list(factura.cartera.pagos.order_by("fecha_pago", "id"))
        for p in pagos:
            p.documento = p.numero_recibo if p.forma_pago == "EFECTIVO" else p.referencia
            p.es_nota_credito = p.forma_pago == "NOTA_CREDITO"

    return {
        "factura": factura,
        "carga": carga,
        "pasos": pasos,
        "entregas": entregas,
        "productos": productos,
        "pagos": pagos,
        "total_pagado": _r(sum((p.monto for p in pagos), CERO)),
    }


def _cant(valor):
    texto = f"{valor:,.3f}".rstrip("0").rstrip(".")
    return texto or "0"


# ---------------------------------------------------------------------------
# Listado por ruta y fechas
# ---------------------------------------------------------------------------


def periodo_por_defecto():
    hasta = timezone.localdate()
    return hasta - timedelta(days=DIAS_POR_DEFECTO - 1), hasta


def listado(desde, hasta, ruta_id=None, vendedor_id=None, estado=""):
    """
    Facturas emitidas entre `desde` y `hasta` (incluidas). La ruta es la de
    la carga en la que va o en la que se visitó la factura; con una ruta
    elegida no salen las facturas que todavía no tienen carga.
    Devuelve (facturas, resumen, recortado).
    """
    qs = _base().filter(fecha_emision__gte=desde, fecha_emision__lte=hasta)
    if ruta_id:
        qs = qs.filter(Q(carga__ruta_id=ruta_id) | Q(entregas__carga__ruta_id=ruta_id)).distinct()
    if vendedor_id:
        qs = qs.filter(vendedor_id=vendedor_id)
    qs = qs.order_by("-fecha_emision", "-numero_factura")

    filas = list(qs[: MAX_FILAS_LISTADO + 1])
    recortado = len(filas) > MAX_FILAS_LISTADO
    filas = filas[:MAX_FILAS_LISTADO]
    for f in filas:
        anotar(f)

    _, codigos = FILTROS_ESTADO.get(estado, FILTROS_ESTADO[""])
    if codigos:
        filas = [f for f in filas if f.estado in codigos]

    activas = [f for f in filas if f.estado != "ANULADA"]
    pendientes = {"POR_ENTREGAR", "EN_RUTA", "SIN_VISITAR", "PARCIAL"}
    resumen = {
        "facturas": len(activas),
        "monto": _r(sum((f.total for f in activas), CERO)),
        "entregadas": sum(1 for f in activas if f.estado == "ENTREGADA"),
        "con_devolucion": sum(1 for f in activas if f.estado == "CON_DEVOLUCION"),
        "devuelto": _r(sum((f.valor_devuelto for f in activas), CERO)),
        "pendientes": sum(1 for f in activas if f.estado in pendientes),
        "sin_carga": sum(1 for f in activas if f.estado in pendientes and f.carga_id is None),
        "no_entregadas": sum(1 for f in activas if f.estado == "NO_ENTREGADA"),
        "saldo": _r(sum((f.saldo or CERO for f in activas), CERO)),
        "anuladas": len(filas) - len(activas),
    }
    return filas, resumen, recortado


def ruta_de(factura):
    if factura.carga_id and factura.carga.ruta_id:
        return factura.carga.ruta.nombre
    return ""


def excel_listado(facturas, titulo):
    """Archivo .xlsx (bytes) con el listado, para descargar."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Facturas"
    hoja.append([titulo])
    hoja["A1"].font = Font(bold=True, size=13)
    hoja.append([])
    encabezados = [
        "Factura", "Emitida", "Código cliente", "Cliente", "Sucursal", "Municipio",
        "Canal", "Vendedor", "Modalidad", "Ruta", "Carga", "Repartidor", "Estado",
        "Total", "Entregado", "Devuelto", "Saldo",
    ]
    hoja.append(encabezados)
    relleno = PatternFill("solid", fgColor="417690")
    for celda in hoja[3]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")
    for f in facturas:
        carga = f.carga
        hoja.append(
            [
                f.numero_factura,
                f.fecha_emision,
                f.cliente.codigo,
                f.cliente.nombre,
                f.cliente.sucursal,
                f.cliente.municipio,
                f.cliente.canal,
                f.vendedor.nombre if f.vendedor_id else "",
                f.modalidad_pago.nombre,
                ruta_de(f),
                carga.codigo_reporte if carga else "Sin carga",
                carga.conductor.nombre if carga and carga.conductor_id else "",
                f.estado_nombre,
                float(f.total),
                float(f.valor_entregado),
                float(f.valor_devuelto),
                float(f.saldo) if f.saldo is not None else None,
            ]
        )
    for fila in hoja.iter_rows(min_row=4):
        fila[1].number_format = "DD/MM/YYYY"
        for celda in fila[13:17]:
            celda.number_format = "#,##0.00"
    anchos = [12, 11, 13, 32, 16, 16, 14, 20, 16, 18, 12, 20, 16, 13, 13, 13, 13]
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A4"
    hoja.auto_filter.ref = f"A3:{get_column_letter(len(encabezados))}{hoja.max_row}"

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
