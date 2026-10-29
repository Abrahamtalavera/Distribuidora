"""
Cambio #25: importa el catálogo de sucursales (direcciones de entrega) de
los clientes, desde el maestro de clientes de SAP.

A diferencia de importar_excel (facturas, core/importador.py), este SÍ
actualiza los registros que ya existen: es un catálogo que cambia de vez en
cuando (direcciones corregidas, sucursales nuevas), no un historial de
transacciones que no se deba tocar. La clave es (código de cliente, nombre
de la sucursal); se puede volver a subir el mismo archivo o uno más
reciente sin duplicar nada.
"""

from django.db import transaction

from core.models import Cliente, Sucursal

COLUMNAS_ESPERADAS = [
    "Código SN",
    "Nombre SN",
    "Cod Grupo",
    "Grupo",
    "Nombre de dirección",
    "Departamento",
    "Territorio",
    "Municipio",
    "Direccion",
    "Canal",
    "EstadoCliente",
]


class ImportacionError(Exception):
    pass


def _leer_filas(archivo):
    import openpyxl

    wb = openpyxl.load_workbook(archivo, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    faltantes = [c for c in COLUMNAS_ESPERADAS if c not in headers]
    if faltantes:
        raise ImportacionError(
            "El Excel no tiene las columnas esperadas. Faltan: " + ", ".join(faltantes)
        )
    col_idx = {h: i for i, h in enumerate(headers) if h}

    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col_idx["Código SN"]] is None:
            continue  # fila vacía de relleno al final del archivo
        filas.append({h: row[col_idx[h]] for h in COLUMNAS_ESPERADAS})
    return filas


def _texto(valor):
    return str(valor).strip() if valor is not None else ""


def _cod_grupo(valor):
    try:
        return int(valor)
    except (TypeError, ValueError):
        return None


@transaction.atomic
def importar_excel(archivo, nombre_archivo=""):
    """
    Procesa el Excel y crea o actualiza las Sucursales (y los Clientes que
    todavía no existan). Devuelve un resumen (dict) con contadores para
    mostrar al usuario.
    """
    filas = _leer_filas(archivo)

    clientes_creados_antes = Cliente.objects.count()
    sucursales_creadas = 0
    sucursales_actualizadas = 0
    filas_sin_codigo_o_nombre = 0

    cache_clientes = {}

    for fila in filas:
        codigo = _texto(fila["Código SN"])
        nombre_sucursal = _texto(fila["Nombre de dirección"])
        if not codigo or not nombre_sucursal:
            filas_sin_codigo_o_nombre += 1
            continue

        cliente = cache_clientes.get(codigo)
        if cliente is None:
            cliente, _creado = Cliente.objects.get_or_create(
                codigo=codigo,
                defaults={"nombre": _texto(fila["Nombre SN"]) or codigo},
            )
            cache_clientes[codigo] = cliente

        activo = _texto(fila["EstadoCliente"]).lower() != "inactivo"
        _sucursal, creada = Sucursal.objects.update_or_create(
            cliente=cliente,
            nombre=nombre_sucursal,
            defaults={
                "direccion": _texto(fila["Direccion"]),
                "departamento": _texto(fila["Departamento"]),
                "municipio": _texto(fila["Municipio"]),
                "territorio": _texto(fila["Territorio"]),
                "canal": _texto(fila["Canal"]),
                "cod_grupo": _cod_grupo(fila["Cod Grupo"]),
                "grupo": _texto(fila["Grupo"]),
                "activo": activo,
            },
        )
        if creada:
            sucursales_creadas += 1
        else:
            sucursales_actualizadas += 1

    return {
        "filas_en_excel": len(filas),
        "sucursales_creadas": sucursales_creadas,
        "sucursales_actualizadas": sucursales_actualizadas,
        "clientes_nuevos": Cliente.objects.count() - clientes_creados_antes,
        "filas_omitidas": filas_sin_codigo_o_nombre,
    }
